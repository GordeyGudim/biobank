"""Выгрузка всех таблиц базы в Excel для просмотра (по листу на таблицу)."""
import sqlite3, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

db = sys.argv[1] if len(sys.argv) > 1 else "vietnam_fish.db"
out = sys.argv[2] if len(sys.argv) > 2 else "vietnam_fish_tables.xlsx"
con = sqlite3.connect(db)
sheets = [("Проблемы (data_issue)", "SELECT * FROM data_issue ORDER BY category, issue_id"),
          ("Особи — сводно", "SELECT * FROM v_specimen"),
          ("expedition", "SELECT * FROM expedition"), ("event", "SELECT * FROM event"),
          ("haul", "SELECT * FROM haul"), ("haul_point", "SELECT * FROM haul_point"),
          ("taxon", "SELECT * FROM taxon"), ("specimen", "SELECT * FROM specimen"),
          ("sample", "SELECT * FROM sample"), ("catch_record", "SELECT * FROM catch_record")]
wb = Workbook(); wb.remove(wb.active)
for title, q in sheets:
    cur = con.execute(q)
    ws = wb.create_sheet(title[:31])
    cols = [d[0] for d in cur.description]
    ws.append(cols)
    for row in cur: ws.append(list(row))
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7")
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(cols, 1):
        width = max([len(str(col))] + [len(str(v.value)) for v in ws[get_column_letter(i)][1:200] if v.value is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(8, width + 2), 60)
wb.save(out)
print("saved", out)
