"""Чтение Excel-книги в промежуточные объекты (dataclass) — без базы данных.

Здесь только «что написано в таблице»: значения разобраны парсерами, но правила
(какое это место, какая дата у рынка, объединять ли точки) применяет importer.py.

Колонки ищутся по тексту шапки, а не по номеру: на разных листах они сдвинуты.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from biobank.parsing import (
    City,
    clean_comment,
    is_catch_comment,
    is_text_number,
    normalize_space,
    parse_city,
    parse_date,
    parse_gear,
    parse_label_range,
    parse_labs,
    parse_number,
    parse_time,
    parse_trawl_point,
    sheet_capture_method,
)

# Район работ (дельта): всё, что вне этих границ, — подозрительная координата
LAT_RANGE = (8.0, 12.0)
LON_RANGE = (104.0, 108.0)

LEGEND_SHEET_PREFIX = "Условные"


class ExcelReadError(Exception):
    """Ошибка чтения книги с указанием листа и ячейки."""


@dataclass
class Issue:
    """Спорное значение → будущая запись data_issue."""

    category: str
    description: str
    sheet: str | None = None
    cell: str | None = None
    object_label: str | None = None
    raw_value: str | None = None
    table_name: str | None = None
    record_id: int | None = None


@dataclass
class WaterPoint:
    """Строка таблицы тралений: «старт N», «финиш N» или «точка лова рыбака»."""

    point_type: str  # 'старт' | 'финиш' | 'точка'
    cell: str  # ячейка с подписью строки, напр. 'L3'
    measured_time: str | None = None
    lat: float | None = None
    lon: float | None = None
    coord_source: str | None = None
    garmin_wp: str | None = None
    depth_m: float | None = None
    o2_pct: float | None = None
    ppm: float | None = None
    conductivity_ms: float | None = None
    conductivity_us: float | None = None
    ph: float | None = None
    water_temp_c: float | None = None
    notes: list[str] = field(default_factory=list)


@dataclass
class TrawlData:
    trawl_no: int
    points: dict[str, WaterPoint] = field(default_factory=dict)
    time_start: str | None = None
    time_end: str | None = None
    duration_min: float | None = None
    speeds: list[float] = field(default_factory=list)
    intermediate_depth_m: float | None = None
    catch_comments: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)


@dataclass
class CoordinateEntry:
    """Подписанные координаты на листе: «Координаты», «2 рыбак (138-142)» и т.п."""

    text: str
    cell: str
    lat: float
    lon: float
    label_range: tuple[int, int] | None  # номера особей, пойманных в этой точке


@dataclass
class SheetData:
    """Один лист = один выезд."""

    name: str
    capture_method: str
    header_text: str | None
    event_date: str | None
    city: City | None
    gear: tuple[str | None, str | None] | None = None
    trawls: dict[int, TrawlData] = field(default_factory=dict)
    fisher_point: WaterPoint | None = None  # строка «точка лова рыбака»
    fisher_point_comment: str | None = None
    coordinates: list[CoordinateEntry] = field(default_factory=list)
    header_notes: list[str] = field(default_factory=list)
    dna_labs: list[str] = field(default_factory=list)
    smear_labs: list[str] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)


@dataclass
class ExpeditionData:
    region: str  # имя листа-разделителя: 'Донг-хап', 'Виньлонг', 'Кантхо'
    sheets: list[SheetData] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Вспомогательный класс: чтение ячеек листа с учётом предупреждений
# ---------------------------------------------------------------------------


class _SheetReader:
    """Читает ячейки одного листа и собирает предупреждения в список issues."""

    def __init__(self, ws: Worksheet, issues: list[Issue]):
        self.ws = ws
        self.name = ws.title
        self.issues = issues
        self.text_numbers: list[str] = []  # адреса ячеек, где число записано текстом

    def value(self, row: int, col: int | None):
        return None if col is None else self.ws.cell(row, col).value

    def comment(self, row: int, col: int) -> str | None:
        cell = self.ws.cell(row, col)
        return clean_comment(cell.comment.text) if cell.comment else None

    def number(self, row: int, col: int | None, obj: str, table: str) -> float | None:
        raw = self.value(row, col)
        result = parse_number(raw)
        cell = self.ws.cell(row, col).coordinate if col else None
        if is_text_number(raw):
            self.text_numbers.append(cell)
        if result.warning:
            self.issues.append(
                Issue("формат", result.warning, self.name, cell, obj, str(raw), table)
            )
        return result.value

    def time(self, row: int, col: int | None, obj: str, table: str) -> str | None:
        raw = self.value(row, col)
        result = parse_time(raw)
        if result.warning:
            cell = self.ws.cell(row, col).coordinate
            self.issues.append(
                Issue("формат", result.warning, self.name, cell, obj, str(raw), table)
            )
        return result.value


def _header_key(value) -> str | None:
    return normalize_space(value.lower()) if isinstance(value, str) else None


# ---------------------------------------------------------------------------
# Поиск шапок
# ---------------------------------------------------------------------------


def find_specimen_header(ws: Worksheet) -> int | None:
    """Номер строки с шапкой «№ пробы» (1-я или 2-я строка)."""
    for row in range(1, 6):
        key = _header_key(ws.cell(row, 1).value)
        if key and key.startswith("№ пробы"):
            return row
    return None


def find_trawl_header(ws: Worksheet) -> tuple[int, dict[str, int]] | None:
    """Строка с «№ траления» и словарь «поле → номер колонки»."""
    for row in range(1, 9):
        for col in range(1, ws.max_column + 1):
            if _header_key(ws.cell(row, col).value) != "№ траления":
                continue
            cols: dict[str, int] = {"trawl": col}
            for c in range(col + 1, ws.max_column + 1):
                key = _header_key(ws.cell(row, c).value)
                if not key:
                    continue
                if "garmin" in key and "результата" in key:
                    cols["wp"] = c
                elif key.startswith("o2"):
                    cols["o2"] = c
                elif key == "ppm":
                    cols["ppm"] = c
                elif key == "ms":
                    cols["ms"] = c
                elif key == "us":
                    cols["us"] = c
                elif key == "ph":
                    cols["ph"] = c
                elif key.startswith("n ("):
                    cols["lat"] = c
                elif key.startswith("e ("):
                    cols["lon"] = c
                elif key.startswith("промежуточная"):
                    cols["idepth"] = c
                elif key.startswith("глубина"):
                    cols["depth"] = c
                elif key.startswith("t старта"):
                    cols["t_start"] = c
                elif key.startswith("t финиша"):
                    cols["t_end"] = c
                elif key.startswith("t траления"):
                    cols["duration"] = c
                elif key.startswith("t воды"):
                    cols["water_t"] = c
                elif key.startswith("v"):
                    cols["speed"] = c
            return row, cols
    return None


def _coord_source(ws: Worksheet, header_row: int, cols: dict[str, int]) -> str | None:
    key = _header_key(ws.cell(header_row, cols["lat"]).value) if "lat" in cols else None
    if key is None:
        return None
    return "Google Maps" if "google" in key else "Garmin"


# ---------------------------------------------------------------------------
# Разбор частей листа
# ---------------------------------------------------------------------------


def _read_header_comments(r: _SheetReader, sheet: SheetData, rows: set[int]) -> None:
    """Комментарии к ячейкам-шапкам: сеть, лаборатории, общие заметки."""
    for row in sorted(rows):
        for col in range(1, r.ws.max_column + 1):
            text = r.comment(row, col)
            if not text:
                continue
            head = _header_key(r.value(row, col)) or ""
            if "ячеи" in text:
                sheet.gear = parse_gear(text)
            elif head == "dna" and parse_labs(text):
                sheet.dna_labs = parse_labs(text)
            elif head.startswith("мазок") and parse_labs(text):
                sheet.smear_labs = parse_labs(text)
            elif "шприц" in text:
                continue  # методика взятия крови, одинаковая на всех листах
            else:
                sheet.header_notes.append(f"{r.value(row, col)}: {text}")


def _read_coordinates(r: _SheetReader, sheet: SheetData) -> None:
    """Подписанные координаты: «Координаты», «Координаты 1 (155-162)», «2 рыбак (138-142)»."""
    for row in r.ws.iter_rows():
        for cell in row:
            text = cell.value
            if not isinstance(text, str):
                continue
            low = text.lower()
            if not (low.startswith("координаты") or "рыбак (" in low):
                continue
            lat_raw = r.value(cell.row, cell.column + 1)
            lon_raw = r.value(cell.row, cell.column + 2)
            if (lat_raw, lon_raw) == ("N", "E"):  # значения строкой ниже (Аквахозяйство)
                lat_raw = r.value(cell.row + 1, cell.column + 1)
                lon_raw = r.value(cell.row + 1, cell.column + 2)
            lat, lon = parse_number(lat_raw).value, parse_number(lon_raw).value
            if lat is None or lon is None:
                continue
            sheet.coordinates.append(
                CoordinateEntry(
                    normalize_space(text), cell.coordinate, lat, lon, parse_label_range(text)
                )
            )
            _check_area(r, lat, lon, cell.coordinate, normalize_space(text), "sampling_site")


def _check_area(r: _SheetReader, lat, lon, cell: str, obj: str, table: str) -> None:
    if lat is not None and not LAT_RANGE[0] <= lat <= LAT_RANGE[1]:
        r.issues.append(
            Issue("координаты", "Широта вне района работ", r.name, cell, obj, str(lat), table)
        )
    if lon is not None and not LON_RANGE[0] <= lon <= LON_RANGE[1]:
        r.issues.append(
            Issue("координаты", "Долгота вне района работ", r.name, cell, obj, str(lon), table)
        )


def _read_trawls(r: _SheetReader, sheet: SheetData, header_row: int, cols: dict[str, int]):
    source = _coord_source(r.ws, header_row, cols)
    time_cols = {cols.get("t_start"), cols.get("t_end"), cols.get("duration")}

    for row in range(header_row + 1, r.ws.max_row + 1):
        parsed = parse_trawl_point(r.value(row, cols["trawl"]))
        if parsed is None:
            continue
        trawl_no, point_type = parsed
        label_cell = r.ws.cell(row, cols["trawl"]).coordinate
        obj = f"траление {trawl_no} {point_type}" if trawl_no else "точка лова рыбака"
        table = "water_measurement"

        point = WaterPoint(point_type, label_cell, coord_source=source)
        wp = r.value(row, cols.get("wp"))
        point.garmin_wp = None if wp is None else str(wp).strip()
        point.lat = r.number(row, cols.get("lat"), obj, table)
        point.lon = r.number(row, cols.get("lon"), obj, table)
        point.depth_m = r.number(row, cols.get("depth"), obj, table)
        point.o2_pct = r.number(row, cols.get("o2"), obj, table)
        point.ppm = r.number(row, cols.get("ppm"), obj, table)
        point.conductivity_ms = r.number(row, cols.get("ms"), obj, table)
        point.conductivity_us = r.number(row, cols.get("us"), obj, table)
        point.ph = r.number(row, cols.get("ph"), obj, table)
        point.water_temp_c = r.number(row, cols.get("water_t"), obj, table)
        lat_cell = r.ws.cell(row, cols["lat"]).coordinate if "lat" in cols else label_cell
        _check_area(r, point.lat, point.lon, lat_cell, obj, table)

        if point_type == "точка":
            sheet.fisher_point = point
            sheet.fisher_point_comment = r.comment(row, cols["trawl"])
            _read_point_comments(r, row, cols, point, None, header_row, time_cols)
            continue

        trawl = sheet.trawls.setdefault(trawl_no, TrawlData(trawl_no))
        if point_type in trawl.points:
            trawl.issues.append(
                Issue("связь", f"Строка «{point_type}» траления повторяется", r.name,
                      label_cell, obj, None, "trawling")
            )  # fmt: skip
        trawl.points[point_type] = point

        if point_type == "старт":
            trawl.time_start = r.time(row, cols.get("t_start"), obj, "trawling")
            point.measured_time = trawl.time_start
            trawl.intermediate_depth_m = r.number(row, cols.get("idepth"), obj, "trawling")
        else:
            trawl.time_end = r.time(row, cols.get("t_end"), obj, "trawling")
            point.measured_time = trawl.time_end
        duration = r.number(row, cols.get("duration"), obj, "trawling")
        if duration is not None:
            trawl.duration_min = duration
        speed = r.number(row, cols.get("speed"), obj, "trawling")
        if speed is not None:
            trawl.speeds.append(speed)

        _read_point_comments(r, row, cols, point, trawl, header_row, time_cols)

    for trawl in sheet.trawls.values():
        missing = {"старт", "финиш"} - trawl.points.keys()
        if missing:
            trawl.issues.append(
                Issue("связь", f"У траления нет строки: {', '.join(sorted(missing))}", r.name,
                      None, f"траление {trawl.trawl_no}", None, "trawling")
            )  # fmt: skip


def _read_point_comments(r, row, cols, point, trawl, header_row, time_cols) -> None:
    """Комментарии в строке траления: об улове, о береге, о параллельных замерах."""
    for col in range(cols["trawl"], r.ws.max_column + 1):
        text = r.comment(row, col)
        if not text:
            continue
        if col == cols["trawl"]:
            if trawl is None:
                point.notes.append(text)
            elif is_catch_comment(text):
                trawl.catch_comments.append(text)
            else:
                trawl.notes.append(text)
        elif col in time_cols and trawl is not None:
            trawl.notes.append(text)
        else:
            head = r.value(header_row, col)
            head = normalize_space(str(head)) if head else r.ws.cell(row, col).coordinate
            point.notes.append(f"{head}: {text}")


def read_sheet(ws: Worksheet) -> SheetData:
    """Прочитать лист-выезд (без особей — они добавятся на этапе 4)."""
    issues: list[Issue] = []
    r = _SheetReader(ws, issues)

    spec_header = find_specimen_header(ws)
    if spec_header is None:
        raise ExcelReadError(f"Лист «{ws.title}»: не найдена шапка «№ пробы» в колонке A")
    header_text = ws.cell(1, 1).value if spec_header > 1 else None
    header_text = normalize_space(header_text) if isinstance(header_text, str) else None

    sheet = SheetData(
        name=ws.title,
        capture_method=sheet_capture_method(ws.title),
        header_text=header_text,
        event_date=parse_date(header_text),
        city=parse_city(header_text),
        issues=issues,
    )

    trawl_header = find_trawl_header(ws)
    header_rows = {spec_header}
    if trawl_header:
        header_rows.add(trawl_header[0])
        _read_trawls(r, sheet, *trawl_header)
    _read_header_comments(r, sheet, header_rows)
    _read_coordinates(r, sheet)

    if r.text_numbers:
        cells = r.text_numbers
        issues.append(
            Issue(
                "формат",
                f"Чисел, записанных текстом (запятая, пробел, точка в конце): {len(cells)}; "
                "при импорте прочитаны как числа",
                ws.title,
                f"{cells[0]}…{cells[-1]}" if len(cells) > 1 else cells[0],
                table_name="sampling_event",
            )
        )
    return sheet


def _is_empty_sheet(ws: Worksheet) -> bool:
    return all(cell.value is None for row in ws.iter_rows() for cell in row)


def read_workbook(path: Path | str) -> list[ExpeditionData]:
    """Прочитать всю книгу: пустые листы-разделители задают экспедиции."""
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as error:  # битый файл, не Excel, нет доступа
        raise ExcelReadError(f"Не удалось открыть книгу {path}: {error}") from error
    expeditions: list[ExpeditionData] = []
    for ws in wb.worksheets:
        if ws.title.startswith(LEGEND_SHEET_PREFIX):
            continue
        if _is_empty_sheet(ws):
            expeditions.append(ExpeditionData(region=ws.title.strip()))
            continue
        if not expeditions:
            raise ExcelReadError(f"Лист «{ws.title}» стоит раньше первого листа-разделителя")
        try:
            expeditions[-1].sheets.append(read_sheet(ws))
        except ExcelReadError:
            raise
        except Exception as error:  # неожиданная ошибка — добавить имя листа
            raise ExcelReadError(f"Лист «{ws.title}»: {error}") from error
    return expeditions
