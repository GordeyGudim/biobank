"""Выгрузка базы в Excel (xlsx) или CSV.

Порядок листов: журнал проблем (data_issue) → сводные таблицы → все таблицы базы.

Учётные записи сайта (app_user) выгружаются без данных для входа: только номер,
имя для подписи, ссылка на человека, активна ли, когда создана (EXPORT_COLUMNS).
Номер нужен, чтобы понять created_by в других таблицах; хеш пароля и логин в файл,
который пересылают коллегам, попадать не должны.
"""

from __future__ import annotations

import csv
import re
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from biobank.db import list_tables
from biobank.reports import SUMMARY_QUERIES

ISSUES_SHEET = "Проблемы (data_issue)"
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")

# порядок таблиц в выгрузке: от справочников к данным; остальные — по алфавиту в конце
TABLE_ORDER = (
    "region", "city", "sampling_site", "capture_method", "gear", "species", "lab", "person",
    "expedition", "expedition_member", "sampling_event", "trawling", "water_measurement",
    "specimen", "species_identification", "sample", "sample_shipment", "analysis",
    "catch_record", "app_user", "user_role", "change_log",
)  # fmt: skip

# Таблицы, у которых выгружаются только перечисленные колонки (остальные — все).
# Список «что можно», а не «что нельзя»: новая колонка app_user (например, токен
# сброса пароля) в выгрузку не попадёт, пока её не добавят сюда явно.
EXPORT_COLUMNS = {
    "app_user": ("user_id", "display_name", "person_id", "is_active", "created_at"),
}


def export_sheets(conn: sqlite3.Connection) -> list[tuple[str, str]]:
    """Список (имя листа, SQL) в порядке выгрузки."""
    sheets = [(ISSUES_SHEET, "SELECT * FROM data_issue ORDER BY resolved, category, issue_id")]
    sheets += list(SUMMARY_QUERIES.items())
    tables = [t for t in list_tables(conn) if t != "data_issue"]
    ordered = [t for t in TABLE_ORDER if t in tables] + sorted(set(tables) - set(TABLE_ORDER))
    # имена таблиц — из самой базы (sqlite_master), колонки — из EXPORT_COLUMNS в коде;
    # данные пользователя в текст запроса не попадают
    for table in ordered:
        columns = ", ".join(EXPORT_COLUMNS.get(table, ("*",)))
        sheets.append((table, f"SELECT {columns} FROM {table}"))
    return sheets


def _sheet_title(name: str) -> str:
    """Имя листа Excel: не длиннее 31 символа и без символов []:*?/\\."""
    return re.sub(r"[\[\]:*?/\\]", "_", name)[:31]


def export_xlsx(conn: sqlite3.Connection, out_path: Path) -> list[tuple[str, int]]:
    """Одна книга, по листу на таблицу. Возвращает [(лист, строк)]."""
    wb = Workbook()
    wb.remove(wb.active)
    written = []
    for name, sql in export_sheets(conn):
        cursor = conn.execute(sql)
        columns = [d[0] for d in cursor.description]
        ws = wb.create_sheet(_sheet_title(name))
        ws.append(columns)
        rows = 0
        for row in cursor:
            ws.append(list(row))
            rows += 1
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
        ws.freeze_panes = "A2"  # шапка не уезжает при прокрутке
        ws.auto_filter.ref = ws.dimensions  # фильтры в шапке
        for i, column in enumerate(columns, start=1):
            letter = get_column_letter(i)
            sample = [len(str(c.value)) for c in ws[letter][1:200] if c.value is not None]
            ws.column_dimensions[letter].width = min(max([8, len(column) + 2, *sample]) + 1, 60)
        written.append((name, rows))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return written


def export_csv(conn: sqlite3.Connection, out_dir: Path) -> list[tuple[str, int]]:
    """По CSV-файлу на таблицу.

    Кодировка utf-8-sig — UTF-8 с меткой BOM в начале файла: по ней Excel понимает,
    что это UTF-8, и показывает кириллицу правильно.
    """
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for index, (name, sql) in enumerate(export_sheets(conn), start=1):
        cursor = conn.execute(sql)
        safe_name = re.sub(r"\W+", "_", name).strip("_")
        file_name = f"{index:02d}_{safe_name}.csv"
        with open(out_dir / file_name, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow([d[0] for d in cursor.description])
            rows = 0
            for row in cursor:
                writer.writerow(row)
                rows += 1
        written.append((file_name, rows))
    return written
