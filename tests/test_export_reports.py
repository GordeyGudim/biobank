"""Тесты этапа 6: выгрузка xlsx/csv, отчёты, команда sql."""

import csv

import openpyxl
import pytest
from conftest import SOURCE, count

from biobank.__main__ import main
from biobank.db import connect_read_only
from biobank.export import ISSUES_SHEET, export_csv, export_xlsx
from biobank.reports import REPORTS, render_report
from biobank.tables import run_query

pytestmark = pytest.mark.skipif(not SOURCE.exists(), reason="нет data/source.xlsx")


@pytest.fixture
def ro(db_path):
    """Подключение только на чтение к импортированной базе."""
    conn = connect_read_only(db_path)
    yield conn
    conn.close()


# ---------------------------------------------------------------------------
# Выгрузка
# ---------------------------------------------------------------------------


def test_xlsx_sheets(ro, tmp_path):
    out = tmp_path / "out.xlsx"
    written = dict(export_xlsx(ro, out))
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames[0] == ISSUES_SHEET  # первый лист — журнал проблем
    assert wb.sheetnames[1:4] == ["Особи — сводно", "Образцы — сводно", "Вода — сводно"]
    assert "specimen" in wb.sheetnames and "water_measurement" in wb.sheetnames
    assert written["specimen"] == 462
    assert written["Особи — сводно"] == 462
    assert written["sample"] == count(ro, "SELECT count(*) FROM sample")
    assert wb["specimen"].max_row == 462 + 1  # + строка шапки


def test_xlsx_values(ro, tmp_path):
    out = tmp_path / "out.xlsx"
    export_xlsx(ro, out)
    ws = openpyxl.load_workbook(out)["Особи — сводно"]
    header = [c.value for c in ws[1]]
    rows = {r[0]: dict(zip(header, r, strict=True)) for r in ws.iter_rows(2, values_only=True)}
    assert rows["155 pc"]["SL, см"] == pytest.approx(29.9)
    assert rows["138 pc"]["способ получения"] == "рыбак"
    assert rows["26 ph"]["способ получения"] == "аквахозяйство"


def test_csv_files(ro, tmp_path):
    written = export_csv(ro, tmp_path / "csv")
    names = [name for name, _ in written]
    assert names[0] == "01_Проблемы_data_issue.csv"
    path = tmp_path / "csv" / names[1]
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")  # BOM — чтобы Excel понял UTF-8
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    assert rows[0][0] == "№ пробы"
    assert len(rows) == 462 + 1


# ---------------------------------------------------------------------------
# Отчёты
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("name", list(REPORTS))
def test_every_report_renders(ro, name):
    text = render_report(ro, name)
    assert "┌" in text and "(строк:" in text


def test_species_report_totals(ro):
    _, rows = run_query(ro, REPORTS["species"][1][1])
    assert sum(r[1] for r in rows) == 462


def test_samples_report_totals(ro):
    _, rows = run_query(ro, REPORTS["samples"][0][1])
    assert sum(r[2] for r in rows) == count(ro, "SELECT count(*) FROM sample")


def test_sites_report_shared_fisher_site(ro):
    columns, rows = run_query(ro, REPORTS["sites"][0][1])
    site = next(dict(zip(columns, r, strict=True)) for r in rows if r[3] == 9.5187059)
    assert site["выездов"] == 2
    assert site["даты"] == "2026-08-21,2026-08-22"


# ---------------------------------------------------------------------------
# Команда sql — только чтение
# ---------------------------------------------------------------------------


def test_sql_select(db_path, capsys):
    assert main(["--db", str(db_path), "sql", "SELECT count(*) AS n FROM specimen"]) == 0
    assert "462" in capsys.readouterr().out


@pytest.mark.parametrize(
    "query",
    ["DELETE FROM species", "UPDATE specimen SET tl_cm = 1", "DROP TABLE lab"],
)
def test_sql_refuses_to_write(db_path, capsys, query):
    assert main(["--db", str(db_path), "sql", query]) == 1
    assert "только читает данные" in capsys.readouterr().err
    conn = connect_read_only(db_path)
    try:
        assert count(conn, "SELECT count(*) FROM species") == 8
        assert count(conn, "SELECT count(*) FROM lab") == 4
    finally:
        conn.close()


def test_sql_error_is_friendly(db_path, capsys):
    assert main(["--db", str(db_path), "sql", "SELECT * FROM nope"]) == 1
    assert "Ошибка SQL" in capsys.readouterr().err
