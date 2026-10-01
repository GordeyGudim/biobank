"""Приёмочные тесты импорта (этап 4): особи, определения, образцы, лаборатории, прилов."""

import openpyxl
import pytest
from conftest import SOURCE, count

from biobank.excel_reader import find_specimen_columns, find_specimen_header

pytestmark = pytest.mark.skipif(not SOURCE.exists(), reason="нет data/source.xlsx")


def samples(db, specimen_label):
    rows = db.execute(
        "SELECT sa.sample_type, sa.organ, sa.label FROM sample sa "
        "JOIN specimen s USING (specimen_id) WHERE s.label = ? ORDER BY sa.sample_id",
        (specimen_label,),
    )
    return [tuple(r) for r in rows]


# ---------------------------------------------------------------------------
# Особи
# ---------------------------------------------------------------------------


def test_specimen_counts(db):
    assert count(db, "SELECT count(*) FROM specimen") == 462
    assert count(db, "SELECT count(*) FROM specimen WHERE label LIKE 'ХХ%'") == 1
    assert count(db, "SELECT count(*) FROM specimen WHERE event_id IS NULL") == 0
    assert count(db, "SELECT count(DISTINCT label) FROM specimen") == 462


def test_trawling_not_guessed(db):
    assert count(db, "SELECT count(*) FROM specimen WHERE trawling_id IS NOT NULL") == 0


def test_label_normalized(db):
    raw = dict(db.execute("SELECT label, label_raw FROM specimen").fetchall())
    assert raw["155 pc"] == "155 pc"  # в Excel «155  pc» (два пробела), пробелы схлопнуты
    assert raw["1 e"] == "1е"  # кириллическая «е» в исходнике
    assert raw["3 pc"] == "3 p.c."
    assert raw["107"] == "107.0"


def test_integrity(db):
    assert db.execute("PRAGMA foreign_key_check").fetchall() == []
    assert db.execute("PRAGMA integrity_check").fetchone()[0] == "ok"


def test_measurements_match_excel(db):
    """Сверка с исходником: каждое числовое TL и m в Excel совпадает с базой (0 расхождений)."""
    stored = {
        (r["source_sheet"], r["source_row"]): (r["tl_cm"], r["mass_g"])
        for r in db.execute(
            "SELECT e.source_sheet, s.source_row, s.tl_cm, s.mass_g "
            "FROM specimen s JOIN sampling_event e USING (event_id)"
        )
    }
    wb = openpyxl.load_workbook(SOURCE)
    checked, mismatches = 0, []
    for ws in wb.worksheets:
        header = find_specimen_header(ws)
        if header is None:
            continue
        cols = find_specimen_columns(ws, header)
        for row in range(header + 1, ws.max_row + 1):
            if (ws.title, row) not in stored:
                continue
            for key, index in (("tl", 0), ("mass", 1)):
                value = ws.cell(row, cols[key]).value
                if isinstance(value, (int, float)):
                    checked += 1
                    if stored[(ws.title, row)][index] != pytest.approx(float(value)):
                        mismatches.append((ws.title, row, key, value))
    assert checked > 800
    assert mismatches == []


def test_excel_date_restored(db):
    assert count(db, "SELECT sl_cm FROM specimen WHERE label = '155 pc'") == pytest.approx(29.9)


def test_whole_specimen_has_no_measurements(db):
    row = db.execute(
        "SELECT s.tl_cm, s.mass_g FROM specimen s JOIN sample sa USING (specimen_id) "
        "WHERE sa.sample_type = 'целая особь' AND sa.raw_value = 'целая рыба' LIMIT 1"
    ).fetchone()
    assert tuple(row) == (None, None)


def test_dead_and_sex(db):
    assert count(db, "SELECT is_dead FROM specimen WHERE label = '116 pc'") == 1  # «особи мертвые»
    assert count(db, "SELECT is_dead FROM specimen WHERE label = '1'") == 1  # «из погибшей особи»
    assert count(db, "SELECT sex FROM specimen WHERE label = '26 pc'") == "самка"


# ---------------------------------------------------------------------------
# Место поимки внутри выезда
# ---------------------------------------------------------------------------


def capture(db, label):
    return db.execute(
        "SELECT m.name, st.lat, st.lon FROM specimen s "
        "LEFT JOIN capture_method m ON m.capture_method_id = s.capture_method_id "
        "LEFT JOIN sampling_site st ON st.site_id = s.capture_site_id WHERE s.label = ?",
        (label,),
    ).fetchone()


@pytest.mark.parametrize("label", ["138 pc", "140 pc", "142 pc"])
def test_fisher_specimens_on_shared_site(db, label):
    method, lat, lon = capture(db, label)
    assert method == "рыбак"
    assert (lat, lon) == (pytest.approx(9.5187059), pytest.approx(106.2082146))


def test_fisher_site_has_events_and_specimens(db):
    site_id = count(db, "SELECT site_id FROM sampling_site WHERE abs(lat - 9.5187059) < 1e-6")
    assert count(db, "SELECT count(*) FROM sampling_event WHERE site_id = ?", site_id) == 2
    assert count(db, "SELECT count(*) FROM specimen WHERE capture_site_id = ?", site_id) == 5


def test_capture_only_when_different(db):
    assert tuple(capture(db, "118 pc"))[0] == "рыбак"
    assert tuple(capture(db, "64 pc"))[0] == "рыбак"  # «далее рыба, пойманная рыбаком»
    assert tuple(capture(db, "69 pc"))[0] == "рыбак"
    assert tuple(capture(db, "63 pc")) == (None, None, None)
    assert tuple(capture(db, "155 pc")) == (None, None, None)  # место = место выезда
    method, lat, _ = capture(db, "163 pc")  # «Координаты 2» на листе рыбака
    assert method is None and lat == pytest.approx(9.60547)


# ---------------------------------------------------------------------------
# Определение вида
# ---------------------------------------------------------------------------


def test_every_specimen_identified(db):
    assert (
        count(
            db,
            "SELECT count(*) FROM specimen s WHERE NOT EXISTS "
            "(SELECT 1 FROM species_identification i WHERE i.specimen_id = s.specimen_id)",
        )
        == 0
    )


def test_identification_matches_current_species(db):
    assert (
        count(
            db,
            "SELECT count(*) FROM specimen s JOIN species_identification i USING (specimen_id) "
            "WHERE s.species_id IS NOT i.species_id",
        )
        == 0
    )


@pytest.mark.parametrize(
    "raw, confidence",
    [("пред. Pangasius elongatus", "предп."), ("Pangasius sp", "до рода"), ("?", "не определён")],
)
def test_confidence(db, raw, confidence):
    assert (
        count(db, "SELECT DISTINCT confidence FROM species_identification WHERE raw_text = ?", raw)
        == confidence
    )


def test_species_comment_in_identification_notes(db):
    notes = [r[0] for r in db.execute("SELECT notes FROM species_identification") if r[0]]
    assert any("возможно pangasius macronema" in n for n in notes)
    assert any("предп. kremfi" in n for n in notes)


# ---------------------------------------------------------------------------
# Образцы и лаборатории
# ---------------------------------------------------------------------------


def test_histology_by_organ(db):
    assert ("гистология", "печень", "26 ph a") in samples(db, "26 ph")
    assert ("гистология", "почки", "26 ph б") in samples(db, "26 ph")
    histo_11 = [s for s in samples(db, "11 pc") if s[0] == "гистология"]
    assert histo_11 == [
        ("гистология", "печень", "11 pc a"),
        ("гистология", "почки", "11 pc b"),
        ("гистология", "икра", "11 pc с"),
    ]


def test_histology_copy_paste_kept_as_written(db):
    assert ("гистология", "почки", "2 pc b") in samples(db, "3 pc")
    assert count(db, "SELECT count(*) FROM data_issue WHERE description LIKE '%копипаст%'") == 2


def test_swapped_columns_point2(db):
    assert samples(db, "1 pc")[1:] == [
        ("гистология", "печень", "1 pc"),
        ("мазок крови", None, "1 pc"),
    ]


def test_shrimp_samples(db):
    assert samples(db, "1 mr") == [
        ("гистология", "мышцы", "1 mr"),
        ("гистология", "жабры", "1 mr"),
        ("другое", "кишечник, гепатопанкреас, карапакс", "1 mr"),
    ]


def test_smear_preservative(db):
    assert (
        count(
            db,
            "SELECT count(*) FROM sample WHERE sample_type = 'мазок крови' "
            "AND preservative IS NOT 'этанол 96% (фиксация)'",
        )
        == 0
    )


def test_no_smear_for_dead_or_small(db):
    assert [s for s in samples(db, "116 pc") if s[0] == "мазок крови"] == []
    assert [s for s in samples(db, "17 pc") if s[0] == "мазок крови"] == []


def labs_of(db, label, sample_type):
    rows = db.execute(
        "SELECT l.name FROM sample_shipment sh JOIN lab l USING (lab_id) "
        "JOIN sample sa USING (sample_id) JOIN specimen s USING (specimen_id) "
        "WHERE s.label = ? AND sa.sample_type = ? ORDER BY l.lab_id",
        (label, sample_type),
    )
    return [r[0] for r in rows]


def test_shipments_2026(db):
    assert labs_of(db, "143 pc", "ДНК") == ["Россия", "Тропцентр", "Кантхо (секвенирование)"]
    assert labs_of(db, "143 pc", "мазок крови") == ["Россия", "Тропцентр"]
    assert labs_of(db, "26 ph", "ДНК") == []  # 2025: куда отправлено — не записано


def test_missing_shipment_reported_not_guessed(db):
    # на листах «Рыбак 1/2 точка» нет комментария к шапке DNA — лаборатории не угадываем
    assert labs_of(db, "73 pc", "ДНК") == []
    sheets = [
        r[0]
        for r in db.execute(
            "SELECT sheet_name FROM data_issue WHERE table_name = 'sample_shipment' "
            "AND description LIKE '%ДНК%' ORDER BY issue_id"
        )
    ]
    assert "Рыбак 1 точка" in sheets and "Рыбак 2 точка" in sheets


@pytest.mark.parametrize("label", ["70 pc", "104 pc", "166 pc"])
def test_vietnam_company(db, label):
    assert "Вьетнамская компания (ген. анализ)" in labs_of(db, label, "ДНК")


# ---------------------------------------------------------------------------
# Прилов
# ---------------------------------------------------------------------------


def test_catch_from_trawl_comments(db):
    rows = db.execute(
        "SELECT c.taxon_text, c.count_n FROM catch_record c "
        "JOIN trawling t USING (trawling_id) JOIN sampling_event e ON e.event_id = t.event_id "
        "WHERE e.source_sheet = 'Точка 5 2026' AND t.trawl_no = 6"
    ).fetchall()
    # «5 pangasius sp, scaptophagus; купили у другого рыбака 2 плотосуса» — купленное не прилов
    assert [tuple(r) for r in rows] == [("Pangasius", 5)]


def test_catch_notes_under_table(db):
    row = db.execute(
        "SELECT c.taxon_text, c.count_n, t.trawl_no FROM catch_record c "
        "LEFT JOIN trawling t USING (trawling_id) WHERE c.raw_text LIKE '2 креветки%'"
    ).fetchone()
    assert tuple(row) == ("креветки", 2, 4)
    assert (
        count(db, "SELECT mass_g FROM catch_record WHERE raw_text = 'm креветок: 1754 г.'") == 1754
    )
